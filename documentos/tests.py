from datetime import date
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, RequestFactory, SimpleTestCase
from pypdf import PdfReader

from django.core.files.uploadedfile import SimpleUploadedFile

from .importacion import extraer_asunto, extraer_asunto_desde_nombre, extraer_destinatario
from .models import Documento, Adjunto, NotaRecibida
from .scanner import compilar_imagenes_a_pdf
from .views import generar_pdf
from .formato import leer_docx_con_listas, cuerpo_html
from .context_processors import sidebar_context


class ExtraccionTests(SimpleTestCase):
    def test_listas_y_tablas_en_orden_del_word(self):
        from docx import Document
        archivo = Document()
        archivo.add_paragraph('Materiales')
        archivo.add_paragraph('Cable de red', style='List Bullet')
        archivo.add_paragraph('Revisar equipo', style='List Number')
        archivo.add_paragraph('Instalar equipo', style='List Number')
        tabla = archivo.add_table(rows=2, cols=2)
        tabla.cell(0, 0).text = 'Cantidad'
        tabla.cell(0, 1).text = 'Descripción'
        tabla.cell(1, 0).text = '2'
        tabla.cell(1, 1).text = 'Switch de 24 bocas POE'
        archivo.add_paragraph('Sin otro particular')
        buffer = BytesIO()
        archivo.save(buffer)
        buffer.seek(0)
        texto = leer_docx_con_listas(buffer)
        self.assertIn('• Cable de red', texto)
        self.assertIn('1. Revisar equipo\n2. Instalar equipo', texto)
        self.assertIn('• 2 — Switch de 24 bocas POE\nSin otro particular', texto)

    def test_tablas_markdown_docx(self):
        from docx import Document
        archivo = Document()
        archivo.add_paragraph('Materiales solicitados')
        tabla = archivo.add_table(rows=2, cols=2)
        tabla.cell(0, 0).text = 'Cantidad'
        tabla.cell(0, 1).text = 'Descripción'
        tabla.cell(1, 0).text = '2'
        tabla.cell(1, 1).text = 'Switch POE'
        buffer = BytesIO()
        archivo.save(buffer)
        buffer.seek(0)
        texto = leer_docx_con_listas(buffer, tablas_markdown=True)
        self.assertIn('| Cantidad | Descripción |', texto)
        self.assertIn('| :---: | :--- |', texto)
        self.assertIn('| 2 | Switch POE |', texto)

    def test_listas_html_escapado(self):
        html = cuerpo_html('• Cable <especial>\n2. Instalar\n<script>alert(1)</script>')
        self.assertTrue('&bull;&nbsp;' in html or '•&nbsp;' in html)
        self.assertIn('2.&nbsp;', html)
        self.assertIn('&lt;especial&gt;', html)
        self.assertNotIn('<script>', html)

    def test_nombres_compuestos(self):
        self.assertEqual(extraer_asunto_desde_nombre('notaCompraAlcohol.docx'), 'Nota Compra Alcohol')
        self.assertEqual(extraer_asunto_desde_nombre('informe0002-PCmesaEntrada.docx'), 'Informe 0002 PC mesa Entrada')
        self.assertEqual(extraer_asunto_desde_nombre('actaEntregaopdsanantonio.docx'), 'Acta entrega OPD San Antonio')

    def test_asunto_explicito(self):
        self.assertEqual(extraer_asunto('Fecha\nASUNTO: Compra de equipos\nCuerpo', 'notaCompra.docx'), 'Compra de equipos')

    def test_destinatario_completo(self):
        bloque = 'Al Sr. Secretario\nDe Niñez, Adolescencia y Familia\nDel Ministerio de Desarrollo Humano\nDr. Sergio Jaime A. Vidaurre'
        self.assertEqual(extraer_destinatario('Fecha\n'+bloque+'\nSu / Despacho:\nTengo el agrado de dirigirme'), bloque)

    def test_para_no_incluye_remitente(self):
        self.assertEqual(extraer_destinatario('PARA: Coordinación de OPD\nDE: Informática\nASUNTO: Equipo'), 'PARA: Coordinación de OPD')

    def test_sin_destinatario(self):
        self.assertEqual(extraer_destinatario('Fecha\nORIGINAL\nRecibí del Área de Informática'), '')


class PDFTests(TestCase):
    def test_membrete_en_todas_las_paginas_y_destinatario_largo(self):
        user = User.objects.create_user('prueba')
        destinatario = 'Al Sr. Secretario\n' + 'Organismo provincial completo ' * 8 + '\nNombre Apellido Final'
        doc = Documento.objects.create(tipo='NOTA', fecha=date(2026, 9, 10), asunto='Compra de equipos',
            remitente='Informática', destinatario=destinatario, cuerpo=('Contenido de prueba. ' * 30 + '\n\n') * 30, creado_por=user)
        doc.refresh_from_db()
        self.assertEqual(doc.destinatario, destinatario)
        request = RequestFactory().get('/')
        request.user = user
        response = generar_pdf(request, doc.pk)
        self.assertEqual(response.status_code, 200)
        pdf = PdfReader(BytesIO(response.content))
        self.assertGreater(len(pdf.pages), 1)
        self.assertIn('NOMBRE APELLIDO FINAL', pdf.pages[0].extract_text())
        for page in pdf.pages:
            self.assertGreater(len(page.images), 0)
            self.assertIn('Familia', page.extract_text())
            # Los párrafos largos deben dividirse en líneas, no dibujarse fuera del A4.
            lineas = page.extract_text().splitlines()
            self.assertLess(max(map(len, lineas)), 180)


class PlantillasSNAFTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('admin_snaf', password='password123')
        self.client.force_login(self.user)

    def test_cargar_plantilla_insumos_get(self):
        response = self.client.get('/documentos/nuevo/?plantilla=insumos')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Solicitud urgente de adquisición de diversos insumos')
        self.assertContains(response, 'Dra. MARTA IRIARTE')

    def test_cargar_plantilla_acta_laptop_get(self):
        response = self.client.get('/documentos/nuevo/?plantilla=laptop')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Acta de entrega y asignación de equipamiento portátil')

    def test_crear_documento_desde_plantilla(self):
        response = self.client.post('/documentos/nuevo/', {
            'tipo': 'NOTA',
            'fecha': '2026-09-22',
            'remitente': 'Ing. Carlos Mendoza',
            'destinatario': 'Dra. Marta Iriarte',
            'destinatario_cargo': 'A LA SECRETARIA',
            'destinatario_nombre': 'DRA. MARTA IRIARTE',
            'asunto': 'Solicitud urgente de insumos TI',
            'cuerpo': 'Solicitud formal de discos SSD y estabilizadores para SNAF.',
            'estado': 'EMITIDO',
            'accion_guardar': 'emitir'
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        doc = Documento.objects.filter(asunto='Solicitud urgente de insumos TI').first()
        self.assertIsNotNone(doc)
        self.assertEqual(doc.tipo, 'NOTA')
        self.assertEqual(doc.estado, 'EMITIDO')

    def test_preformatos_dependen_del_tipo_seleccionado(self):
        response = self.client.get('/documentos/nuevo/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'actualizarPreformatosPorTipo')
        self.assertContains(response, 'PREFORMATOS_POR_TIPO')
        self.assertContains(response, 'preformato-banner-titulo')
        self.assertContains(response, 'preformato-select')

    def test_pdf_informe_encabezado_y_doble_firma(self):
        doc = Documento.objects.create(
            tipo='INFORME',
            fecha=date(2026, 8, 3),
            asunto='Estado de impresoras HP LaserJet P1005',
            remitente='Darío Joaquín Orquera',
            cuerpo='Revisión de impresoras.',
            creado_por=self.user
        )
        response = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(response.status_code, 200)
        pdf = PdfReader(BytesIO(response.content))
        texto = '\n'.join([p.extract_text() for p in pdf.pages])
        self.assertIn('SECRETARÍA DE NIÑEZ, ADOLESCENCIA Y FAMILIA', texto)
        self.assertIn('Departamento de Informática', texto)
        self.assertIn('INFORME TÉCNICO', texto)
        self.assertIn('Estado de impresoras HP LaserJet P1005', texto)
        self.assertIn('Darío Joaquín Orquera', texto)
        self.assertIn('Alexander Magaña', texto)
        self.assertIn('Departamento Informática', texto)
        self.assertIn('Sin otro particular, se eleva el presente informe', texto)

    def test_detalle_informe_doble_firma(self):
        doc = Documento.objects.create(
            tipo='INFORME',
            fecha=date(2026, 8, 3),
            asunto='Estado de impresoras HP LaserJet P1005',
            remitente='Darío Joaquín Orquera',
            cuerpo='Revisión de impresoras.',
            creado_por=self.user
        )
        response = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Darío Joaquín Orquera')
        self.assertContains(response, 'Alexander Magaña')
        self.assertContains(response, 'Departamento Informática')
        self.assertContains(response, 'INFORME TÉCNICO')
    def test_pdf_nota_estandar_una_sola_pagina_con_firmas_al_pie(self):
        cuerpo = (
            "Por medio del presente, se informa la necesidad de intervención técnica por parte del CRC, "
            "a fin de realizar la revisión y reparación de dos (2) equipos de impresión pertenecientes "
            "a las áreas de Coordinación de OPD y Dirección Provincial de Niñez, Adolescencia y Familia.\n"
            "Detalle de equipos y fallas detectadas:\n"
            "1. Equipo: HP Smart Tank 530\n"
            "Falla detectada: El equipo presenta un mal funcionamiento en el sistema de alimentación de tinta.\n"
            "2. Equipo: Xerox Phaser 3020\n"
            "Falla detectada: El equipo presenta líneas blancas verticales constantes en las impresiones.\n"
            "Solicitud: Se solicita la evaluación técnica de ambos equipos y la ejecución de las tareas de reparación necesarias."
        )
        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 20),
            asunto='Pedido de reparacion de impresoras a juventud',
            remitente='Informática',
            destinatario="AL DIRECTOR DE\nJUVENTUD\nDR. DARIO NICOLAS D'ANTUENE",
            cuerpo=cuerpo,
            creado_por=self.user
        )
        response = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(response.status_code, 200)
        pdf = PdfReader(BytesIO(response.content))
        self.assertEqual(len(pdf.pages), 1)
        y_coords = []
        def vis(text, cm, tm, font_dict, font_size):
            if 'Alexander' in text or 'Darío' in text or 'Dario' in text:
                y_coords.append(cm[5])
        pdf.pages[0].extract_text(visitor_text=vis)
        self.assertTrue(len(y_coords) > 0)
        self.assertLess(max(y_coords), 160)

    def test_acta_duplicada_pdf_y_html(self):
        from .models import ItemActa
        doc = Documento.objects.create(
            tipo='ACTA',
            fecha=date(2026, 8, 3),
            asunto='Acta de entrega de equipamiento',
            remitente='Darío Joaquín Orquera',
            destinatario='Coordinación de Informática',
            cuerpo='Recibí del Área de Informática los bienes que se detallan a continuación.-',
            creado_por=self.user
        )
        ItemActa.objects.create(documento=doc, cantidad=1, descripcion='Monitor 19 Samsung', numero_serie='SN12345')
        ItemActa.objects.create(documento=doc, cantidad=2, descripcion='Teclado USB Genius')

        # 1. Test vista web HTML
        response_html = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(response_html.status_code, 200)
        self.assertContains(response_html, 'ORIGINAL')
        self.assertContains(response_html, 'DUPLICADO')
        self.assertContains(response_html, 'Cortar por la línea de puntos')
        self.assertContains(response_html, 'RECIBE CONFORME:')
        self.assertContains(response_html, 'AUTORIZA ENTREGA:')
        self.assertContains(response_html, 'Monitor 19 Samsung')
        self.assertNotContains(response_html, 'sheet-institution-header')

        # 2. Test generación PDF (debe caber exactamente en 1 página con Original y Duplicado)
        response_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(response_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(response_pdf.content))
        self.assertEqual(len(pdf.pages), 1)
        texto = pdf.pages[0].extract_text()
        self.assertIn('ORIGINAL', texto)
        self.assertIn('DUPLICADO', texto)
        self.assertIn('RECIBE CONFORME:', texto)
        self.assertIn('AUTORIZA ENTREGA:', texto)
        self.assertIn('Monitor 19 Samsung', texto)
        self.assertIn('Cortar por', texto)

    def test_crear_documento_post_invalido_renderiza_sin_error(self):
        # Enviar POST con datos incompletos (sin remitente ni asunto)
        response = self.client.post('/documentos/nuevo/', {
            'tipo': 'ACTA',
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Redactar Nuevo Documento')


class RemitenteAutocompletarTests(TestCase):
    def test_remitente_por_defecto_para_nota_e_informe(self):
        user = User.objects.create_user(
            username='tecnico1',
            first_name='Juan',
            last_name='Pérez',
            password='pass'
        )
        self.client.force_login(user)
        
        # NOTA por defecto
        response_nota = self.client.get('/documentos/nuevo/')
        self.assertEqual(response_nota.status_code, 200)
        self.assertEqual(response_nota.context['form'].initial['remitente'], 'Área de Informática')
        self.assertContains(response_nota, 'value="Área de Informática"')

        # INFORME
        response_inf = self.client.get('/documentos/nuevo/?tipo=INFORME')
        self.assertEqual(response_inf.status_code, 200)
        self.assertEqual(response_inf.context['form'].initial['remitente'], 'Área de Informática')
        self.assertContains(response_inf, 'value="Área de Informática"')

    def test_autocompletar_remitente_acta_con_nombre_completo(self):
        user = User.objects.create_user(
            username='tecnico1',
            first_name='Juan',
            last_name='Pérez',
            password='pass'
        )
        self.client.force_login(user)
        response = self.client.get('/documentos/nuevo/?tipo=ACTA')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['form'].initial['remitente'], 'Juan Pérez')
        self.assertContains(response, 'value="Juan Pérez"')

    def test_autocompletar_remitente_acta_sin_nombre_completo_usa_username(self):
        user = User.objects.create_user(
            username='admin_soporte',
            first_name='',
            last_name='',
            password='pass'
        )
        self.client.force_login(user)
        response = self.client.get('/documentos/nuevo/?tipo=ACTA')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['form'].initial['remitente'], 'admin_soporte')
        self.assertContains(response, 'value="admin_soporte"')

    def test_sidebar_context_nombre_usuario(self):
        rf = RequestFactory()

        # Con nombre y apellido
        user_con_nombre = User.objects.create_user(
            username='usuario_con_nombre',
            first_name='Carlos',
            last_name='Gómez',
            password='pass'
        )
        req1 = rf.get('/')
        req1.user = user_con_nombre
        ctx1 = sidebar_context(req1)
        self.assertEqual(ctx1['nombre_usuario'], 'Carlos Gómez')

        # Sin nombre ni apellido (fallback al username)
        user_sin_nombre = User.objects.create_user(
            username='operador1',
            first_name='',
            last_name='',
            password='pass'
        )
        req2 = rf.get('/')
        req2.user = user_sin_nombre
        ctx2 = sidebar_context(req2)
        self.assertEqual(ctx2['nombre_usuario'], 'operador1')


class ActaReceptorTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='tecnico_acta',
            first_name='Darío Joaquín',
            last_name='Orquera',
            password='password123'
        )
        self.client.force_login(self.user)

    def test_acta_con_datos_receptor_en_html_y_pdf(self):
        from .models import ItemActa
        doc = Documento.objects.create(
            tipo='ACTA',
            fecha=date(2026, 9, 23),
            asunto='Entrega de periféricos a Mesa de Entrada',
            remitente='Darío Joaquín Orquera',
            destinatario='Mesa de Entrada',
            receptor_nombre='Juan Carlos Pérez',
            receptor_dni='35.123.456',
            cuerpo='Recibí del Área de Informática los bienes detallados a continuación.',
            creado_por=self.user
        )
        ItemActa.objects.create(documento=doc, cantidad=1, descripcion='Monitor Samsung 24', numero_serie='SN998877')

        # 1. Vista Web HTML (detalle): Original y Duplicado deben mostrar los datos del receptor y sector
        resp_html = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(resp_html.status_code, 200)
        self.assertContains(resp_html, 'Juan Carlos Pérez')
        self.assertContains(resp_html, '35.123.456')
        self.assertContains(resp_html, 'DESTINO:')
        self.assertContains(resp_html, 'Mesa de Entrada')
        self.assertContains(resp_html, 'DEPENDENCIA / SECTOR:')
        html_content = resp_html.content.decode('utf-8')
        self.assertGreaterEqual(html_content.count('35.123.456'), 2)
        self.assertGreaterEqual(html_content.count('Mesa de Entrada'), 2)

        # 2. PDF: Original y Duplicado deben contener los datos del receptor y sector destino
        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        self.assertEqual(len(pdf.pages), 1)
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('Juan Carlos Pérez', texto_pdf)
        self.assertIn('35.123.456', texto_pdf)
        self.assertIn('DESTINO:', texto_pdf)
        self.assertIn('MESA DE ENTRADA', texto_pdf)
        self.assertIn('DEPENDENCIA / SECTOR:', texto_pdf)

    def test_acta_sin_datos_receptor_mantiene_lineas_puntos(self):
        doc = Documento.objects.create(
            tipo='ACTA',
            fecha=date(2026, 9, 23),
            asunto='Entrega con receptor manuscrito',
            remitente='Darío Joaquín Orquera',
            receptor_nombre='',
            receptor_dni='',
            destinatario='',
            cuerpo='Recibí del Área de Informática los bienes.',
            creado_por=self.user
        )

        resp_html = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(resp_html.status_code, 200)
        self.assertContains(resp_html, 'ACLARACION:...........................................................DNI...................................')
        self.assertContains(resp_html, 'DEPENDENCIA / SECTOR: ........................................................................................')

        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('ACLARACION:', texto_pdf)
        self.assertIn('DNI', texto_pdf)
        self.assertIn('DEPENDENCIA / SECTOR:', texto_pdf)

    def test_crear_acta_post_con_receptor(self):
        datos = {
            'tipo': 'ACTA',
            'fecha': '2026-09-23',
            'asunto': 'Entrega de notebook',
            'remitente': 'Darío Joaquín Orquera',
            'destinatario': 'Dirección Provincial',
            'receptor_nombre': 'Mariana Sánchez',
            'receptor_dni': '32.456.789',
            'cuerpo': 'Recibí del Área de Informática.',
            'accion_guardar': 'emitir',
            'items-TOTAL_FORMS': '1',
            'items-INITIAL_FORMS': '0',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
            'items-0-cantidad': '1',
            'items-0-descripcion': 'Notebook Lenovo ThinkPad',
            'items-0-numero_serie': 'LNV12345',
            'items-0-condicion': 'NUEVO',
        }
        resp = self.client.post('/documentos/nuevo/', datos)
        self.assertEqual(resp.status_code, 302)
        doc = Documento.objects.get(asunto='Entrega de notebook')
        self.assertEqual(doc.receptor_nombre, 'Mariana Sánchez')
        self.assertEqual(doc.receptor_dni, '32.456.789')
        self.assertEqual(doc.items.count(), 1)

    def test_acta_cuerpo_por_defecto_en_formulario_get(self):
        resp = self.client.get('/documentos/nuevo/?tipo=ACTA')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp.context['form'].initial['cuerpo'],
            'Recibí del Área de Informática los bienes que se detallan a continuación.-'
        )
        self.assertContains(resp, 'Recibí del Área de Informática los bienes que se detallan a continuación.-')

    def test_acta_cuerpo_por_defecto_al_guardar_modelo_sin_cuerpo(self):
        doc = Documento.objects.create(
            tipo='ACTA',
            fecha=date(2026, 9, 23),
            asunto='Acta sin cuerpo explícito',
            remitente='Tester',
            creado_por=self.user
        )
        self.assertEqual(doc.cuerpo, 'Recibí del Área de Informática los bienes que se detallan a continuación.-')

    def test_acta_cuerpo_por_defecto_al_postear_con_cuerpo_vacio(self):
        datos = {
            'tipo': 'ACTA',
            'fecha': '2026-09-23',
            'asunto': 'Acta con cuerpo vacío en POST',
            'remitente': 'Darío Joaquín Orquera',
            'cuerpo': '',
            'accion_guardar': 'emitir',
            'items-TOTAL_FORMS': '0',
            'items-INITIAL_FORMS': '0',
            'items-MIN_NUM_FORMS': '0',
            'items-MAX_NUM_FORMS': '1000',
        }
        resp = self.client.post('/documentos/nuevo/', datos)
        self.assertEqual(resp.status_code, 302)
        doc = Documento.objects.get(asunto='Acta con cuerpo vacío en POST')
        self.assertEqual(doc.cuerpo, 'Recibí del Área de Informática los bienes que se detallan a continuación.-')


class AdjuntosEImagenesTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username='tecnico_adjuntos',
            first_name='Carlos',
            last_name='Mendoza',
            password='password123'
        )
        self.client.force_login(self.user)

    def test_modelo_adjunto_propiedades(self):
        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 23),
            asunto='Nota con adjuntos',
            remitente='Área de Informática',
            creado_por=self.user
        )
        archivo_fake = SimpleUploadedFile('informe_switch.pdf', b'%PDF-1.4 test content', content_type='application/pdf')
        adj = Adjunto.objects.create(
            documento=doc,
            archivo=archivo_fake,
            nombre_original='informe_switch.pdf',
            subido_por=self.user
        )
        self.assertTrue(adj.es_pdf)
        self.assertFalse(adj.es_imagen)
        self.assertEqual(adj.extension, 'pdf')
        self.assertIn('KB', adj.tamano_formateado)

    def test_crear_documento_con_archivos_adjuntos(self):
        f1 = SimpleUploadedFile('foto_placa.png', b'\x89PNG\r\n\x1a\nfake', content_type='image/png')
        f2 = SimpleUploadedFile('presupuesto.pdf', b'%PDF-1.4 fake', content_type='application/pdf')
        datos = {
            'tipo': 'NOTA',
            'fecha': '2026-09-23',
            'asunto': 'Nota con dos adjuntos',
            'remitente': 'Área de Informática',
            'cuerpo': 'Se adjunta la documentación solicitada.',
            'accion_guardar': 'borrador',
            'archivos': [f1, f2],
        }
        resp = self.client.post('/documentos/nuevo/', datos)
        self.assertEqual(resp.status_code, 302)
        doc = Documento.objects.get(asunto='Nota con dos adjuntos')
        self.assertEqual(doc.adjuntos.count(), 2)

        resp_det = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(resp_det.status_code, 200)
        self.assertContains(resp_det, 'foto_placa.png')
        self.assertContains(resp_det, 'presupuesto.pdf')
        self.assertContains(resp_det, 'Archivos Adjuntos y Documentación Respaldatoria')

    def test_subir_imagen_editor_endpoint(self):
        img = SimpleUploadedFile('captura_error.jpg', b'\xff\xd8\xff\xe0 fake jpeg', content_type='image/jpeg')
        resp = self.client.post('/api/subir-imagen/', {'imagen': img})
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertTrue(data['success'])
        self.assertIn('[IMAGEN:', data['tag'])
        self.assertIn('captura_error.jpg', data['nombre'])

    def test_cuerpo_html_renderiza_etiqueta_imagen(self):
        cuerpo = 'Diagnóstico preliminar:\n[IMAGEN: /media/adjuntos/placa.jpg | Placa madre quemada]\nFin del informe.'
        html = cuerpo_html(cuerpo, para_pdf=False)
        self.assertIn('<img src="/media/adjuntos/placa.jpg"', html)
        self.assertIn('Placa madre quemada', html)

        cuerpo_md = '![Esquema de red](/media/adjuntos/red.png)'
        html_md = cuerpo_html(cuerpo_md, para_pdf=False)
        self.assertIn('<img src="/media/adjuntos/red.png"', html_md)
        self.assertIn('Esquema de red', html_md)

    def test_pdf_con_imagen_en_cuerpo_genera_200(self):
        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 23),
            asunto='Nota con imagen en cuerpo',
            remitente='Área de Informática',
            cuerpo='Prueba de imagen:\n[IMAGEN: /media/fake.png | Captura de pantalla]\nContinuación.',
            creado_por=self.user
        )
        resp = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['content-type'], 'application/pdf')
        pdf = PdfReader(BytesIO(resp.content))
        self.assertGreaterEqual(len(pdf.pages), 1)

    def test_lista_documentos_orden_fecha_desc_y_asc(self):
        Documento.objects.create(
            tipo='NOTA',
            fecha=date(2024, 1, 1),
            asunto='Doc Antiguo',
            remitente='Admin',
            cuerpo='Cuerpo',
            creado_por=self.user
        )
        Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 12, 31),
            asunto='Doc Nuevo',
            remitente='Admin',
            cuerpo='Cuerpo',
            creado_por=self.user
        )
        # Mayor a menor (desc)
        resp_desc = self.client.get('/documentos/?orden=fecha&dir=desc')
        self.assertEqual(resp_desc.status_code, 200)
        docs_desc = list(resp_desc.context['documentos'])
        self.assertGreater(docs_desc[0].fecha, docs_desc[-1].fecha)

        # Menor a mayor (asc)
        resp_asc = self.client.get('/documentos/?orden=fecha&dir=asc')
        self.assertEqual(resp_asc.status_code, 200)
        docs_asc = list(resp_asc.context['documentos'])
        self.assertLess(docs_asc[0].fecha, docs_asc[-1].fecha)

    def test_lista_documentos_orden_numero(self):
        resp = self.client.get('/documentos/?orden=numero&dir=desc')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['orden'], 'numero')
        self.assertEqual(resp.context['dir'], 'desc')

    def test_destinatario_cargo_y_nombre_separados_en_nota_e_informe(self):
        doc_nota = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 23),
            asunto='Nota con cargo y nombre separados',
            remitente='Área de Informática',
            destinatario_cargo='a la secretaria de niñez, adolescencia y familia',
            destinatario_nombre='dra. marta iriarte',
            cuerpo='Contenido de la nota solicitando insumos.',
            creado_por=self.user
        )
        doc_nota.refresh_from_db()
        self.assertEqual(doc_nota.destinatario_cargo, 'a la secretaria de niñez, adolescencia y familia')
        self.assertEqual(doc_nota.cargo_display, 'A LA SECRETARIA DE NIÑEZ, ADOLESCENCIA Y FAMILIA')
        self.assertEqual(doc_nota.nombre_display, 'DRA. MARTA IRIARTE')
        self.assertIn('A LA SECRETARIA DE NIÑEZ, ADOLESCENCIA Y FAMILIA', doc_nota.destinatario)
        self.assertIn('DRA. MARTA IRIARTE', doc_nota.destinatario)

        # Verificación en vista detalle HTML
        resp_nota = self.client.get(f'/documentos/{doc_nota.pk}/')
        self.assertEqual(resp_nota.status_code, 200)
        self.assertContains(resp_nota, 'A LA SECRETARIA DE NIÑEZ, ADOLESCENCIA Y FAMILIA')
        self.assertContains(resp_nota, 'DRA. MARTA IRIARTE')
        self.assertContains(resp_nota, 'SU DESPACHO:')

        # Verificación en PDF de Nota
        resp_pdf_nota = self.client.get(f'/documentos/{doc_nota.pk}/pdf/')
        self.assertEqual(resp_pdf_nota.status_code, 200)
        pdf_nota = PdfReader(BytesIO(resp_pdf_nota.content))
        texto_pdf_nota = pdf_nota.pages[0].extract_text()
        self.assertIn('A LA SECRETARIA DE NIÑEZ, ADOLESCENCIA Y FAMILIA', texto_pdf_nota)
        self.assertIn('DRA. MARTA IRIARTE', texto_pdf_nota)
        self.assertIn('SU DESPACHO:', texto_pdf_nota)

        # Creación y verificación para Informe
        doc_inf = Documento.objects.create(
            tipo='INFORME',
            fecha=date(2026, 9, 23),
            asunto='Informe técnico de infraestructura',
            remitente='Área de Informática',
            destinatario_cargo='a la directora provincial de administración y presupuesto',
            destinatario_nombre='cpn leila rivera',
            cuerpo='Contenido técnico detallado del informe.',
            creado_por=self.user
        )
        resp_inf = self.client.get(f'/documentos/{doc_inf.pk}/')
        self.assertEqual(resp_inf.status_code, 200)
        self.assertContains(resp_inf, 'A LA DIRECTORA PROVINCIAL DE ADMINISTRACIÓN Y PRESUPUESTO')
        self.assertContains(resp_inf, 'CPN LEILA RIVERA')
        self.assertContains(resp_inf, 'SU DESPACHO:')

        # Verificación en PDF de Informe
        resp_pdf_inf = self.client.get(f'/documentos/{doc_inf.pk}/pdf/')
        self.assertEqual(resp_pdf_inf.status_code, 200)
        pdf_inf = PdfReader(BytesIO(resp_pdf_inf.content))
        texto_pdf_inf = pdf_inf.pages[0].extract_text()
        self.assertIn('A LA DIRECTORA PROVINCIAL DE ADMINISTRACIÓN Y PRESUPUESTO', texto_pdf_inf)
        self.assertIn('CPN LEILA RIVERA', texto_pdf_inf)
        self.assertIn('SU DESPACHO:', texto_pdf_inf)

    def test_cuerpo_html_renderiza_tabla_markdown_web_y_pdf(self):
        cuerpo_con_tabla = """A continuación se detalla el listado de componentes requeridos:

| Cantidad | Componente / Descripción | Estado |
| :--- | :---: | ---: |
| 02 | **Disco SSD Kingston 240GB** | Nuevo |
| 01 | Memoria RAM DDR4 8GB | Operativo |

Sin otro particular, quedo a su disposición."""

        # 1. Prueba para HTML Web
        html_web = cuerpo_html(cuerpo_con_tabla, para_pdf=False)
        self.assertIn('doc-markdown-table', html_web)
        self.assertIn('<th', html_web)
        self.assertIn('Componente / Descripción', html_web)
        self.assertIn('<strong>Disco SSD Kingston 240GB</strong>', html_web)
        self.assertIn('Memoria RAM DDR4 8GB', html_web)

        # 2. Prueba para PDF
        html_pdf = cuerpo_html(cuerpo_con_tabla, para_pdf=True)
        self.assertIn('tabla-cuerpo-pdf', html_pdf)
        self.assertIn('border: 0.75pt solid', html_pdf)

        # 3. Prueba en documento completo y generación real de PDF
        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 23),
            asunto='Nota con tabla de componentes',
            remitente='Área de Informática',
            destinatario_cargo='A LA SECRETARIA DE NIÑEZ, ADOLESCENCIA Y FAMILIA',
            destinatario_nombre='DRA. MARTA IRIARTE',
            cuerpo=cuerpo_con_tabla,
            creado_por=self.user
        )
        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('Disco SSD Kingston 240GB', texto_pdf)
        self.assertIn('Memoria RAM DDR4 8GB', texto_pdf)

    def test_lista_documentos_filtros_activos_y_desplegable(self):
        # Sin filtros activos
        resp_sin_filtros = self.client.get('/documentos/')
        self.assertEqual(resp_sin_filtros.status_code, 200)
        self.assertEqual(resp_sin_filtros.context['filtros_activos_count'], 0)
        self.assertContains(resp_sin_filtros, 'id="filters-container"')
        self.assertContains(resp_sin_filtros, 'btn-toggle-filters')

        # Con filtros activos (tipo=NOTA y estado=BORRADOR)
        resp_con_filtros = self.client.get('/documentos/?tipo=NOTA&estado=BORRADOR')
        self.assertEqual(resp_con_filtros.status_code, 200)
        self.assertEqual(resp_con_filtros.context['filtros_activos_count'], 2)
        self.assertContains(resp_con_filtros, '2 activos')
        self.assertContains(resp_con_filtros, 'Tipo: NOTA')
        self.assertContains(resp_con_filtros, 'Estado: BORRADOR')


class MesaEntradaNotasRecibidasTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='admin_ti', password='password123', first_name='Juan', last_name='Pérez')
        self.client.login(username='admin_ti', password='password123')

    def test_correlativo_automatico_recibidos(self):
        nota1 = NotaRecibida.objects.create(
            fecha_recepcion=date(2026, 9, 23),
            remitente_origen='Secretaría de Niñez',
            asunto='Solicitud de mantenimiento',
            recibido_por=self.user
        )
        self.assertEqual(nota1.numero_registro, 'REC-001/2026')

        nota2 = NotaRecibida.objects.create(
            fecha_recepcion=date(2026, 9, 23),
            remitente_origen='Dirección de Adultos Mayores',
            asunto='Pedido de mouse',
            recibido_por=self.user
        )
        self.assertEqual(nota2.numero_registro, 'REC-002/2026')

    def test_compilar_imagenes_a_pdf(self):
        from PIL import Image
        img1 = Image.new('RGB', (400, 600), color='white')
        img2 = Image.new('RGB', (400, 600), color='blue')
        pdf_file = compilar_imagenes_a_pdf([img1, img2], nombre_archivo='test.pdf')
        self.assertIsNotNone(pdf_file)
        self.assertTrue(len(pdf_file.read()) > 1000)

    def test_crear_nota_recibida_vista(self):
        from PIL import Image
        img_io = BytesIO()
        Image.new('RGB', (300, 400), color='white').save(img_io, format='JPEG')
        img_io.seek(0)
        archivo_foto = SimpleUploadedFile('foto_pagina1.jpg', img_io.getvalue(), content_type='image/jpeg')

        response = self.client.post('/recibidos/nuevo/', {
            'fecha_recepcion': '2026-09-23',
            'hora_recepcion': '10:30',
            'remitente_origen': 'Defensoría de Niñez',
            'numero_origen': 'Nota N° 88/26',
            'asunto': 'Revisión técnica de impresoras',
            'descripcion': 'La impresora no toma las hojas.',
            'recibido_por': self.user.pk,
            'estado': 'PENDIENTE',
            'imagenes_escaneo': [archivo_foto],
        }, follow=True)

        self.assertEqual(response.status_code, 200)
        nota = NotaRecibida.objects.filter(numero_origen='Nota N° 88/26').first()
        self.assertIsNotNone(nota)
        self.assertEqual(nota.numero_registro, 'REC-001/2026')
        self.assertTrue(bool(nota.archivo_pdf))

        # Descargar PDF compilado
        resp_pdf = self.client.get(f'/recibidos/{nota.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        self.assertEqual(resp_pdf['Content-Type'], 'application/pdf')

    def test_responder_nota_recibida_con_documento(self):
        nota = NotaRecibida.objects.create(
            fecha_recepcion=date(2026, 9, 23),
            remitente_origen='Subsecretaría de Juventud',
            asunto='Pedido de conectividad',
            recibido_por=self.user,
            estado='PENDIENTE'
        )

        # GET crear documento con nota_recibida_id
        resp_get = self.client.get(f'/documentos/nuevo/?nota_recibida_id={nota.pk}')
        self.assertEqual(resp_get.status_code, 200)
        self.assertContains(resp_get, nota.numero_registro)

        # POST crear documento que responde
        resp_post = self.client.post('/documentos/nuevo/', {
            'tipo': 'NOTA',
            'fecha': '2026-09-23',
            'asunto': f'Respuesta a {nota.numero_registro}: Conectividad configurada',
            'remitente': 'Área de Informática',
            'destinatario_cargo': 'SUBSECRETARÍA DE JUVENTUD',
            'destinatario_nombre': 'DR. CARLOS LÓPEZ',
            'cuerpo': 'Se procedió con la instalación solicitada.',
            'estado': 'EMITIDO',
            'nota_recibida_id': nota.pk,
        }, follow=True)

        self.assertEqual(resp_post.status_code, 200)
        nota.refresh_from_db()
        self.assertEqual(nota.estado, 'RESPONDIDO')
        self.assertIsNotNone(nota.documento_respuesta)
        self.assertEqual(nota.documento_respuesta.tipo, 'NOTA')

    def test_sidebar_context_notas_recibidas(self):
        NotaRecibida.objects.create(
            fecha_recepcion=date(2026, 9, 23),
            remitente_origen='Despacho',
            asunto='Prueba de conteo',
            recibido_por=self.user,
            estado='PENDIENTE'
        )
        factory = RequestFactory()
        request = factory.get('/')
        request.user = self.user
        ctx = sidebar_context(request)
        self.assertEqual(ctx['conteo_recibidos_total'], 1)
        self.assertEqual(ctx['conteo_recibidos_pendientes'], 1)

    def test_descargar_pdf_nota_recibida_headers_sameorigin(self):
        from PIL import Image
        img_io = BytesIO()
        Image.new('RGB', (200, 200), color='white').save(img_io, format='JPEG')
        img_io.seek(0)
        archivo_foto = SimpleUploadedFile('test.jpg', img_io.getvalue(), content_type='image/jpeg')

        nota = NotaRecibida.objects.create(
            fecha_recepcion=date(2026, 9, 23),
            remitente_origen='Remitente',
            asunto='Asunto con PDF',
            recibido_por=self.user,
            archivo_pdf=SimpleUploadedFile('documento.pdf', b'%PDF-1.4 test', content_type='application/pdf')
        )

        resp = self.client.get(f'/recibidos/{nota.pk}/pdf/')
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertEqual(resp.headers.get('X-Frame-Options'), 'SAMEORIGIN')

    def test_editar_nota_recibida_eliminar_pdf(self):
        nota = NotaRecibida.objects.create(
            fecha_recepcion=date(2026, 9, 23),
            remitente_origen='Remitente original',
            asunto='Nota para quitar PDF',
            recibido_por=self.user,
            archivo_pdf=SimpleUploadedFile('eliminar_me.pdf', b'%PDF-1.4 dummy', content_type='application/pdf')
        )
        self.assertTrue(bool(nota.archivo_pdf))

        # Enviar POST con eliminar_pdf_actual='1'
        resp = self.client.post(f'/recibidos/{nota.pk}/editar/', {
            'fecha_recepcion': '2026-09-23',
            'remitente_origen': 'Remitente modificado',
            'asunto': 'Nota con PDF quitado',
            'recibido_por': self.user.pk,
            'estado': 'PENDIENTE',
            'eliminar_pdf_actual': '1',
        }, follow=True)

        self.assertEqual(resp.status_code, 200)
        nota.refresh_from_db()
        self.assertFalse(bool(nota.archivo_pdf))
        self.assertEqual(nota.asunto, 'Nota con PDF quitado')

    def test_compilar_imagenes_filtros_camscanner(self):
        from PIL import Image
        img = Image.new('RGB', (400, 500), color='white')

        # Prueba con Magic Color
        pdf_magic = compilar_imagenes_a_pdf([img], nombre_archivo='magic.pdf', filtro='magic_color', autocrop=True)
        self.assertIsNotNone(pdf_magic)

        # Prueba con Blanco y Negro
        pdf_bn = compilar_imagenes_a_pdf([img], nombre_archivo='bn.pdf', filtro='blanco_negro', autocrop=False)
        self.assertIsNotNone(pdf_bn)

        # Prueba con Escala de Grises
        pdf_grises = compilar_imagenes_a_pdf([img], nombre_archivo='grises.pdf', filtro='escala_grises', autocrop=False)
        self.assertIsNotNone(pdf_grises)

    def test_detalle_nota_recibida_renderiza_visor_pdfjs(self):
        nota = NotaRecibida.objects.create(
            fecha_recepcion=date(2026, 9, 23),
            remitente_origen='Test Viewer',
            asunto='Asunto Viewer PDF',
            recibido_por=self.user,
            archivo_pdf=SimpleUploadedFile('test_viewer.pdf', b'%PDF-1.4 test', content_type='application/pdf')
        )
        resp = self.client.get(f'/recibidos/{nota.pk}/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'pdf.min.js')
        self.assertContains(resp, 'pdf-render-canvas')
        self.assertContains(resp, 'pdf-canvas-wrapper')
        self.assertContains(resp, 'pdf-prev-btn')
        self.assertContains(resp, 'pdf-next-btn')

    def test_marcar_como_entregado_y_bloqueo_edicion(self):
        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 29),
            asunto='Nota para entrega',
            remitente='Área de Informática',
            destinatario='Secretaría',
            cuerpo='Contenido de la nota.',
            estado='EMITIDO',
            creado_por=self.user,
        )
        # 1. Marcar como entregado desde la vista de detalle
        resp = self.client.post(f'/documentos/{doc.pk}/', {'accion_entregar': '1'}, follow=True)
        self.assertEqual(resp.status_code, 200)
        doc.refresh_from_db()
        self.assertEqual(doc.estado, 'ENTREGADO')

        # 2. Intentar editar un documento entregado debe ser bloqueado
        resp_edit = self.client.get(f'/documentos/{doc.pk}/editar/', follow=True)
        self.assertEqual(resp_edit.status_code, 200)
        self.assertContains(resp_edit, 'no puede ser modificado')

        # 3. Intentar eliminar un documento entregado debe ser bloqueado
        resp_del = self.client.get(f'/documentos/{doc.pk}/eliminar/', follow=True)
        self.assertEqual(resp_del.status_code, 200)
        self.assertContains(resp_del, 'No se puede eliminar')

    def test_marcar_como_entregado_con_adjuntos_escaneados_y_pdf(self):
        from PIL import Image
        img_io = BytesIO()
        Image.new('RGB', (300, 400), color='white').save(img_io, format='JPEG')
        img_io.seek(0)
        archivo_foto = SimpleUploadedFile('escaneo_pagina1.jpg', img_io.getvalue(), content_type='image/jpeg')

        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 29),
            asunto='Nota con escaneo firmado',
            remitente='Área de Informática',
            destinatario='Secretaría',
            cuerpo='Nota para entregar con firma escaneada.',
            estado='EMITIDO',
            creado_por=self.user,
        )

        resp = self.client.post(f'/documentos/{doc.pk}/', {
            'accion_entregar': '1',
            'imagenes_escaneo': [archivo_foto],
            'filtro_escaneo': 'magic_color',
            'autocrop': '1',
        }, follow=True)

        self.assertEqual(resp.status_code, 200)
        doc.refresh_from_db()
        self.assertEqual(doc.estado, 'ENTREGADO')
        self.assertEqual(doc.adjuntos.count(), 1)
        adj = doc.adjuntos.first()
        self.assertTrue(adj.nombre_original.startswith('FIRMADO_'))
        self.assertTrue(adj.nombre_original.endswith('.pdf'))
        self.assertTrue(adj.es_pdf)


class DashboardYBusquedaGlobalTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='admin_test', password='password123')
        self.client.force_login(self.user)

        self.doc_emitido = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 9, 30),
            asunto='Solicitud de switches de red',
            remitente='Área de Informática',
            destinatario='Dirección General',
            cuerpo='Solicitud formal de equipamiento.',
            estado='EMITIDO',
            creado_por=self.user,
        )

        self.nota_recibida = NotaRecibida.objects.create(
            numero_registro='REC-2026-0099',
            numero_origen='Expte N° 4501/26',
            remitente_origen='Secretaría de Niñez y Familia',
            fecha_recepcion=date(2026, 9, 30),
            asunto='Solicitud de monitores e insumos informáticos',
            descripcion='Nota recibida requiriendo insumos técnicos.',
            estado='PENDIENTE',
            recibido_por=self.user,
        )

    def test_dashboard_incluye_notas_recibidas(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['total_recibidas'], 1)
        self.assertEqual(response.context['recibidas_pendientes'], 1)
        self.assertIn(self.nota_recibida, response.context['recibidas_recientes'])
        self.assertContains(response, 'REC-2026-0099')
        self.assertContains(response, 'Secretaría de Niñez y Familia')

    def test_busqueda_general_incluye_documentos_y_notas_recibidas(self):
        # Búsqueda por término común en ambas o específico de notas recibidas
        response = self.client.get('/documentos/?q=insumos')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['total_recibidas_encontradas'], 1)
        self.assertIn(self.nota_recibida, response.context['notas_recibidas_coincidentes'])
        self.assertContains(response, 'REC-2026-0099')
        self.assertContains(response, 'Mesa de Entrada – Notas Recibidas coincidentes')

    def test_marcar_como_entregado_con_detalle_entrega(self):
        resp = self.client.post(
            f'/documentos/{self.doc_emitido.pk}/',
            {
                'accion_entregar': '1',
                'detalle_entrega': 'Enviado por correo oficial a direccion@snaf.gob.ar con copia a Despacho.',
            },
            follow=True
        )
        self.assertEqual(resp.status_code, 200)
        self.doc_emitido.refresh_from_db()
        self.assertEqual(self.doc_emitido.estado, 'ENTREGADO')
        self.assertEqual(self.doc_emitido.detalle_entrega, 'Enviado por correo oficial a direccion@snaf.gob.ar con copia a Despacho.')
        self.assertContains(resp, 'Constancia y Detalle de Entrega')
        self.assertContains(resp, 'Enviado por correo oficial a direccion@snaf.gob.ar con copia a Despacho.')


class FirmaSelectorTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='admin_firma', password='password123', first_name='Darío', last_name='Orquera')
        self.client.login(username='admin_firma', password='password123')

    def test_formulario_nuevo_contiene_selector_firmas(self):
        resp = self.client.get('/documentos/nuevo/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Firmas al Pie del Documento')
        self.assertContains(resp, 'Darío y Alex (Ambos)')
        self.assertContains(resp, 'Solo Darío')
        self.assertContains(resp, 'Solo Alex')
        self.assertContains(resp, 'id="id_firma"')
        self.assertContains(resp, 'selectFirma')

    def test_crear_documento_con_firma_solo_dario(self):
        datos = {
            'tipo': 'NOTA',
            'fecha': '2026-10-01',
            'asunto': 'Nota firmada solo por Darío',
            'remitente': 'Área de Informática',
            'destinatario_cargo': 'A LA SECRETARIA',
            'destinatario_nombre': 'DRA. MARTA IRIARTE',
            'cuerpo': 'Contenido de prueba.',
            'estado': 'EMITIDO',
            'firma': 'DARIO',
            'accion_guardar': 'emitir',
        }
        resp = self.client.post('/documentos/nuevo/', datos, follow=True)
        self.assertEqual(resp.status_code, 200)
        doc = Documento.objects.get(asunto='Nota firmada solo por Darío')
        self.assertEqual(doc.firma, 'DARIO')

        # Vista Detalle HTML
        resp_det = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(resp_det.status_code, 200)
        self.assertContains(resp_det, 'Darío Joaquín Orquera')
        self.assertNotContains(resp_det, 'Alexander Magaña')

        # Generación PDF
        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('Darío Joaquín Orquera', texto_pdf)
        self.assertNotIn('Alexander Magaña', texto_pdf)

    def test_crear_documento_con_firma_solo_alex(self):
        datos = {
            'tipo': 'INFORME',
            'fecha': '2026-10-01',
            'asunto': 'Informe firmado solo por Alex',
            'remitente': 'Área de Informática',
            'destinatario_cargo': 'A LA DIRECTORA',
            'cuerpo': 'Contenido técnico de prueba.',
            'estado': 'EMITIDO',
            'firma': 'ALEX',
            'accion_guardar': 'emitir',
        }
        resp = self.client.post('/documentos/nuevo/', datos, follow=True)
        self.assertEqual(resp.status_code, 200)
        doc = Documento.objects.get(asunto='Informe firmado solo por Alex')
        self.assertEqual(doc.firma, 'ALEX')

        # Vista Detalle HTML
        resp_det = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(resp_det.status_code, 200)
        self.assertContains(resp_det, 'Alexander Magaña')
        self.assertNotContains(resp_det, 'Darío Joaquín Orquera')

        # Generación PDF
        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('Alexander Magaña', texto_pdf)
        self.assertNotIn('Darío Joaquín Orquera', texto_pdf)

    def test_crear_documento_con_ambas_firmas(self):
        datos = {
            'tipo': 'NOTA',
            'fecha': '2026-10-01',
            'asunto': 'Nota firmada por ambos',
            'remitente': 'Área de Informática',
            'destinatario_cargo': 'A LA SECRETARIA',
            'cuerpo': 'Contenido conjunto.',
            'estado': 'EMITIDO',
            'firma': 'AMBOS',
            'accion_guardar': 'emitir',
        }
        resp = self.client.post('/documentos/nuevo/', datos, follow=True)
        self.assertEqual(resp.status_code, 200)
        doc = Documento.objects.get(asunto='Nota firmada por ambos')
        self.assertEqual(doc.firma, 'AMBOS')

        # Vista Detalle HTML
        resp_det = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(resp_det.status_code, 200)
        self.assertContains(resp_det, 'Darío Joaquín Orquera')
        self.assertContains(resp_det, 'Alexander Magaña')

        # Generación PDF
        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('Darío Joaquín Orquera', texto_pdf)
        self.assertIn('Alexander Magaña', texto_pdf)


class IndentacionTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='user_indent', password='password123')
        self.client.login(username='user_indent', password='password123')

    def test_cuerpo_html_renderiza_sangria_con_espacios_y_tab(self):
        # 4 espacios iniciales (1 tab)
        cuerpo_con_espacios = "    Tengo el agrado de dirigirme a usted..."
        html_espacios = cuerpo_html(cuerpo_con_espacios)
        self.assertIn('&nbsp;&nbsp;&nbsp;&nbsp;Tengo el agrado de dirigirme a usted...', html_espacios)

        # Tabulador inicial \t
        cuerpo_con_tab = "\tPor medio de la presente, informo las novedades..."
        html_tab = cuerpo_html(cuerpo_con_tab)
        self.assertIn('&nbsp;&nbsp;&nbsp;&nbsp;Por medio de la presente, informo las novedades...', html_tab)

        # Doble sangría (8 espacios / 2 tabs)
        cuerpo_doble = "        Sub-párrafo con mayor sangría."
        html_doble = cuerpo_html(cuerpo_doble)
        self.assertIn('&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;Sub-párrafo con mayor sangría.', html_doble)

    def test_documento_con_sangria_pdf_y_html(self):
        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 10, 1),
            asunto='Nota con sangría de párrafo',
            remitente='Área de Informática',
            destinatario_cargo='A LA SECRETARIA DE NIÑEZ',
            cuerpo="    Párrafo 1 con sangría de primera línea.\n\n    Párrafo 2 con sangría de primera línea.",
            creado_por=self.user
        )

        # Vista detalle
        resp_det = self.client.get(f'/documentos/{doc.pk}/')
        self.assertEqual(resp_det.status_code, 200)
        self.assertContains(resp_det, '&nbsp;&nbsp;&nbsp;&nbsp;Párrafo 1 con sangría')

        # PDF
        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('Párrafo 1 con sangría', texto_pdf)
        self.assertIn('Párrafo 2 con sangría', texto_pdf)

    def test_formulario_contiene_boton_sangria(self):
        resp = self.client.get('/documentos/nuevo/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Sangría (Tab)')
        self.assertContains(resp, "formatVisual('indent')")

    def test_formulario_recibido_contiene_buscador_doc_respuesta(self):
        resp = self.client.get('/recibidos/nuevo/')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'id="buscador_doc_respuesta"')
        self.assertContains(resp, 'id="btn_limpiar_busqueda_doc"')
        self.assertContains(resp, 'filtrarDocumentosRespuesta')
        self.assertContains(resp, 'inicializarBuscadorDocRespuesta')
        self.assertContains(resp, 'id="id_documento_respuesta"')

    def test_renglones_en_blanco_cuerpo_html(self):
        # Separación de párrafo estándar (\n\n) -> sin renglón en blanco extra
        html_std = cuerpo_html("Párrafo 1\n\nPárrafo 2")
        self.assertIn('<p style="margin-bottom:6pt;">Párrafo 1</p>', html_std)
        self.assertIn('<p style="margin-bottom:6pt;">Párrafo 2</p>', html_std)
        self.assertNotIn('&nbsp;', html_std)

        # 1 renglón en blanco intencional (\n\n\n) -> un párrafo vacío &nbsp;
        html_1_blank = cuerpo_html("Párrafo 1\n\n\nPárrafo 2")
        self.assertEqual(html_1_blank.count('&nbsp;'), 1)
        self.assertIn('<p style="margin-bottom:6pt;">&nbsp;</p>', html_1_blank)

        # 2 renglones en blanco intencionales (\n\n\n\n) -> dos párrafos vacíos &nbsp;
        html_2_blanks = cuerpo_html("Párrafo 1\n\n\n\nPárrafo 2")
        self.assertEqual(html_2_blanks.count('&nbsp;'), 2)

    def test_documento_con_renglones_en_blanco_pdf(self):
        doc = Documento.objects.create(
            tipo='NOTA',
            fecha=date(2026, 10, 1),
            asunto='Nota con renglones en blanco',
            remitente='Área de Informática',
            destinatario_cargo='A LA SECRETARIA DE NIÑEZ',
            cuerpo="Párrafo 1\n\n\n\nPárrafo 2",
            creado_por=self.user
        )
        resp_pdf = self.client.get(f'/documentos/{doc.pk}/pdf/')
        self.assertEqual(resp_pdf.status_code, 200)
        pdf = PdfReader(BytesIO(resp_pdf.content))
        texto_pdf = pdf.pages[0].extract_text()
        self.assertIn('Párrafo 1', texto_pdf)
        self.assertIn('Párrafo 2', texto_pdf)









