"""
Corrige la fecha de documentos que quedaron con fecha=hoy por no detectarla en texto.
Los re-lee desde el docx buscando la fecha también en tablas y párrafos cortos.
"""
import os, sys, re, django
from pathlib import Path
from datetime import date

sys.path.insert(0, str(Path(__file__).parent))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings')
django.setup()

from docx import Document as DocxDocument
from documentos.models import Documento

CARPETA = r'\\snfserver2\Informatica\notas de pedido'

MESES_ES = {
    'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4,
    'mayo': 5, 'junio': 6, 'julio': 7, 'agosto': 8,
    'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12,
}

def extraer_fecha(texto):
    patron = r'(\d{1,2})\s+(?:de\s+)?(\w+)\s+(?:del?\s+)?(\d{4})'
    for m in re.finditer(patron, texto, re.IGNORECASE):
        dia = int(m.group(1))
        mes_str = m.group(2).lower().strip()
        anio = int(m.group(3))
        mes = MESES_ES.get(mes_str)
        if mes and 1 <= dia <= 31 and 2020 <= anio <= 2030:
            try:
                return date(anio, mes, dia)
            except ValueError:
                pass
    return None


def leer_todo_texto(ruta):
    """Lee párrafos + tablas del docx."""
    doc = DocxDocument(ruta)
    partes = []
    for p in doc.paragraphs:
        if p.text.strip():
            partes.append(p.text.strip())
    for tabla in doc.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                if celda.text.strip():
                    partes.append(celda.text.strip())
    return '\n'.join(partes)


hoy = date.today()
sin_fecha = Documento.objects.filter(fecha=hoy).order_by('numero')
archivos = [f for f in os.listdir(CARPETA) if f.lower().endswith('.docx') and not f.startswith('~')]

corregidos = 0
sin_corregir = []


def clean(s):
    return re.sub(r'[^a-z0-9]', '', s.lower())


for doc in sin_fecha:
    asunto_c = clean(doc.asunto)
    # Buscar archivo mas parecido
    match = None
    best = 0
    for arch in archivos:
        arch_c = clean(arch.replace('.docx', ''))
        comun = sum(1 for a, b in zip(asunto_c, arch_c) if a == b)
        if comun > best:
            best = comun
            match = arch

    if not match or best < 4:
        sin_corregir.append(doc.numero)
        continue

    ruta = os.path.join(CARPETA, match)
    try:
        texto = leer_todo_texto(ruta)
        fecha = extraer_fecha(texto)
        if fecha and fecha != hoy:
            doc.fecha = fecha
            doc.save(update_fields=['fecha'])
            print(f"  OK  {doc.numero} -> {fecha} (desde {match})")
            corregidos += 1
        else:
            sin_corregir.append(f"{doc.numero} ({match})")
    except Exception as e:
        sin_corregir.append(f"{doc.numero} (ERR: {e})")

print(f"\nCorregidos: {corregidos}")
if sin_corregir:
    print(f"Sin corregir ({len(sin_corregir)}): {sin_corregir}")
